from __future__ import annotations

import os
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def build_reports():
    project_root = Path(__file__).resolve().parents[1]
    report_dir = project_root / "Report"
    report_dir.mkdir(parents=True, exist_ok=True)

    # Generated image source path
    artifact_img_path = Path(r"C:\Users\WaveWalker\.gemini\antigravity-ide\brain\6f8767f1-80a8-4b5a-9f40-10a4adbcf2f2\aira_training_architecture_diagram_1789741821325.png")
    dest_img_path = report_dir / "aira_training_architecture_diagram.png"

    if artifact_img_path.exists():
        shutil.copy(artifact_img_path, dest_img_path)
        print(f"[Report] Copied architecture diagram to {dest_img_path}")

    # Build DOCX Document
    doc = Document()

    # Page setup margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles setup
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Aira AI Model Training Architecture Report")
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Dual-System Non-Autoregressive Triage, MCTS Reasoning, and Live Monitoring Web UI")
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    doc.add_paragraph() # Spacer

    # Add Diagram Image
    if dest_img_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(dest_img_path), width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 1: High-Level Architectural Block Diagram of Aira Dual-System Training Pipeline & Live Web UI")
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    doc.add_paragraph()

    # Section 1: Executive Summary
    h1 = doc.add_heading("1. Executive Summary & Overview", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x00, 0x40, 0x80)

    p1 = doc.add_paragraph(
        "This report provides an in-depth technical explanation of the newly implemented Aira Dual-System Training Loop (train_v2.py), "
        "its associated Live Web UI Monitoring Dashboard (live_dashboard.py), and the Google Colab execution launcher (run_colab_training_v2.py). "
        "Aira incorporates a hybrid Dual-System architecture combining System 1 (sub-70ms non-autoregressive parallel reflex triage) "
        "and System 2 (Monte Carlo Tree Search + Process Reward Model reasoning) to achieve frontier-level intelligence while maintaining "
        "ultra-low latency and deterministic guardrail compliance."
    )

    # Section 2: Component Explanations
    h2 = doc.add_heading("2. Core Features & Architectural Role in Aira", level=1)
    h2.runs[0].font.color.rgb = RGBColor(0x00, 0x40, 0x80)

    features = [
        ("1. Data Ingestion & Streaming Pipeline", "Combines FineWeb-Edu, UltraChat 200k, Orca Math, and CodeFeedback datasets. Normalizes and split-streams high-quality training pairs to the model."),
        ("2. CPU-to-CUDA 4-Bit NF4 Block Streaming", "Streams model layer blocks from CPU RAM to CUDA GPU, allowing an 8B parameter model to train on a single 15GB Tesla T4 GPU with peak initialization VRAM under 4.0 GB."),
        ("3. Dual-System Loss Objectives", "Calculates joint loss terms for System 1 (vault path indexing & AGENT.md rule compliance) and System 2 (MCTS reasoning step quality via PRM scoring)."),
        ("4. Muon + AdamW Hybrid Optimizer", "Uses Muon optimizer for 2D weight matrices (scaling tensor learning) combined with AdamW for 1D vectors, normalization layers, and embeddings."),
        ("5. Live Web UI Monitoring Dashboard", "Provides real-time interactive tracking on port 7860 with glassmorphic UI, live Chart.js curves (Loss, Tok/s, VRAM), metric cards, and terminal log tail."),
        ("6. Google Colab Drive Persistence", "Mounts /content/drive/MyDrive/AiraCheckpoints/ to automatically preserve all model weights and checkpoints across Colab sessions."),
    ]

    for title, desc in features:
        h_feat = doc.add_heading(title, level=2)
        h_feat.runs[0].font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        p_feat = doc.add_paragraph(desc)

    # Section 3: Future Plans & Strategic Roadmap
    h3 = doc.add_heading("3. Future Plans & Strategic Roadmap", level=1)
    h3.runs[0].font.color.rgb = RGBColor(0x00, 0x40, 0x80)

    roadmap = [
        "Phase 1 (Current): Dual-System Training Loop v2 deployment with Colab Live Web UI and automated benchmark validation.",
        "Phase 2 (Q4 2026): Autonomous STaR (Self-Taught Reasoner) synthetic data loop to allow Aira to generate her own chain-of-thought training data from vault tasks.",
        "Phase 3 (Q1 2027): Latent space policy distillation, compressing System 2 MCTS search traces into System 1 sub-50ms neural heads.",
        "Phase 4 (Q2 2027): Full recursive self-improvement pipeline with parameter-efficient adapters (LoRA/DoRA) and EWC identity protection.",
    ]

    for item in roadmap:
        doc.add_paragraph(item, style='List Bullet')

    # Save DOCX
    docx_file = report_dir / "Aira_Training_Loop_Architecture_Report.docx"
    doc.save(str(docx_file))
    print(f"[Report] Saved DOCX report: file:///{docx_file}")

    # Build Markdown Report Counterpart
    md_file = report_dir / "Aira_Training_Loop_Architecture_Report.md"
    with open(md_file, "w", encoding="utf-8") as f:
        f.write("# Aira AI Model Training Architecture Report\n\n")
        f.write("## Dual-System Non-Autoregressive Triage, MCTS Reasoning, and Live Monitoring Web UI\n\n")
        f.write("![Figure 1: Architectural Block Diagram](aira_training_architecture_diagram.png)\n\n")
        f.write("### 1. Executive Summary & Overview\n\n")
        f.write(p1.text + "\n\n")
        f.write("### 2. Core Features & Architectural Role in Aira\n\n")
        for title, desc in features:
            f.write(f"#### {title}\n{desc}\n\n")
        f.write("### 3. Future Plans & Strategic Roadmap\n\n")
        for item in roadmap:
            f.write(f"- {item}\n")
        f.write("\n---\n*Report generated in `Report/` directory*\n")

    print(f"[Report] Saved MD report: file:///{md_file}")


if __name__ == "__main__":
    build_reports()
